"""
Excel Exporter
Mengekspor data lowongan kerja ke dalam file Excel (.xlsx) dengan:
- Judul kolom bahasa Indonesia yang ramah (bukan key JSON)
- Kolom kota/area (Jakarta vs Bekasi/Cikarang)
- Bidang spesialisasi (QA/QC, Mikrobiologi, Bioremediasi/WWTP, HACCP/GMP)
- Hyperlink interaktif langsung ke halaman lamaran
- Penataan otomatis lebar kolom & proteksi file lock
"""

import datetime
from pathlib import Path
from typing import List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from scrapers.base_scraper import JobItem

def export_to_excel(jobs: List[JobItem], output_path: Path) -> Path:
    """
    Menyimpan daftar JobItem ke file Excel dengan judul kolom berbahasa Indonesia yang rapi.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Karier Biologi Jabek"
    ws.views.sheetView[0].showGridLines = True

    # Palette Warna Elegan Sains & Biologi (Deep Emerald & Sage)
    HEADER_FILL = PatternFill(start_color="0D5C56", end_color="0D5C56", fill_type="solid")
    HEADER_FONT = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    ZEBRA_FILL = PatternFill(start_color="F2FAF8", end_color="F2FAF8", fill_type="solid")
    REGULAR_FONT = Font(name="Segoe UI", size=10)
    BOLD_FONT = Font(name="Segoe UI", size=10, bold=True)
    LINK_FONT = Font(name="Segoe UI", size=10, color="0066CC", underline="single", bold=True)
    BADGE_FONT = Font(name="Segoe UI", size=9, bold=True, color="0D5C56")

    THIN_SIDE = Side(style='thin', color='D1D5DB')
    BORDER_ALL = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)

    # Judul Kolom Bahasa Indonesia yang Ramah (Bukan Key JSON!)
    headers = [
        "No",
        "Posisi / Judul Lowongan",
        "Nama Perusahaan",
        "Kota / Area",
        "Lokasi Spesifik",
        "Fokus Bidang Keahlian",
        "Kesesuaian Skill (HACCP / GMP / Lab)",
        "Tingkat Pengalaman",
        "Estimasi Gaji",
        "Portal Sumber",
        "Tanggal Posting",
        "Deskripsi & Ringkasan Tugas",
        "Link Pendaftaran Langsung"
    ]

    # Tulis Header
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_ALL
    ws.row_dimensions[1].height = 32

    # Tulis Baris Data
    for idx, job in enumerate(jobs, start=1):
        row_num = idx + 1
        keywords_str = ", ".join(job.matched_keywords) if job.matched_keywords else "-"
        
        # Label area rapi
        if job.city_category == "bekasi":
            area_label = "Bekasi / Cikarang"
        elif job.city_category == "jakarta":
            area_label = "Jakarta"
        else:
            area_label = "Jabodetabek Sekitarnya"

        row_data = [
            idx,
            job.title,
            job.company,
            area_label,
            job.location,
            job.field_category or "Biologi Umum",
            keywords_str,
            job.experience_level or "Fresh Graduate",
            job.salary or "Kompetitif / UMR",
            job.source,
            job.posted_at or "Baru",
            job.description[:250] + ("..." if len(job.description) > 250 else ""),
            "Lamar Pekerjaan Sekarang"
        ]
        ws.append(row_data)

        # Styling per baris
        is_even = (idx % 2 == 0)
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = BORDER_ALL
            cell.alignment = Alignment(vertical="center")

            if is_even:
                cell.fill = ZEBRA_FILL

            # Kolom No, Area, Portal, Tanggal rata tengah
            if col_idx in [1, 4, 10, 11]:
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Judul posisi tebal
            if col_idx == 2:
                cell.font = BOLD_FONT

            # Fokus bidang tebal bernuansa
            if col_idx == 6:
                cell.font = BADGE_FONT

            # Deskripsi wrap text
            if col_idx == 12:
                cell.alignment = Alignment(vertical="center", wrap_text=True)

            # Link lamaran (Hyperlink aktif yang bisa langsung diklik)
            if col_idx == 13 and job.apply_url:
                cell.hyperlink = job.apply_url
                cell.font = LINK_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.row_dimensions[row_num].height = 28

    # Autofit Column Widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        col_idx = col[0].column
        for cell in col:
            val_str = str(cell.value or '')
            if col_idx == 12:  # Kolom deskripsi dibatasi agar tidak terlalu lebar
                max_len = 35
                break
            if col_idx == 13:  # Kolom link
                max_len = 18
                break
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 48)

    # Simpan file dengan penanganan proteksi file lock (PermissionError)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    target_path = output_path

    try:
        wb.save(str(target_path))
    except PermissionError:
        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        fallback_path = output_path.with_name(f"{output_path.stem}_{timestamp}{output_path.suffix}")
        wb.save(str(fallback_path))
        print(f"[Peringatan] File {output_path.name} sedang dibuka di Excel. Tersimpan sebagai file baru: {fallback_path.name}")
        return fallback_path

    return target_path
